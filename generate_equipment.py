"""
관리계획서(CP) -> 설비 일상 점검표 생성기

정책 (사용자 확정 2026-08-27):
    - CP의 '관리기록방법'(AJ열) 값이 '설비 일상 점검표'(= 설비일일점검일지 등 표현 흔들림 허용)인
      행만 위에서부터 순서대로 추출한다.
    - equipment_template.xlsx 의 '1' 시트 레이아웃(No. / 공정단계명 / 공정특성 / 규격 / 1~31일 칸)을
      그대로 쓴다. (다른 시트는 제거하고 이 시트 하나만 남긴다.)
    - No.(A열)      : 추출된 행을 위에서부터 1, 2, 3 ...
    - 공정단계명(B열): CP 공정명(B열). 공정그룹 첫 행에만 값이 있어(병합) 직전 값을 이어받는다.
    - 공정특성(C열)  : CP 공정특성(L열). 검토표시본 주석(🔍/⚠️/❌, [REVIEW], [MISSING])은 잘라낸다.
    - 규격(D열)      : CP 규격하한/Norminal/규격상한/단위(T/U/W/X)를 조합.
    - 표 전체(머리행 ~ 마지막 데이터행, A~AI열)에 all borders(가는 실선 4방향).
    - 시트명 = 설비번호. 예: '135호기' -> '135', 'M135' -> '135', '000' -> '000'.
    - 날짜(1~31)별 점검 칸은 매일 작업자가 채우는 영역이라 비워 둔다.
"""
import re
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from cp_loader import load_cp_workbook

CP_HEADER = {"대상설비": "M3"}

# CP 본문 시작 행 / 열 (다른 generate_*.py 와 동일한 고정 레이아웃 가정)
CP_DATA_START_ROW = 11
CP_COL = {
    "공정단계명": "B",
    "공정특성": "L",
    "규격하한": "T",
    "Norminal": "U",
    "규격상한": "W",
    "단위": "X",
    "관리기록방법": "AJ",
}

# 관리기록방법이 이 값(공백 제거 후)이면 설비 일상 점검표 항목으로 본다.
TARGET_METHODS = {"설비일상점검표", "설비일일점검일지", "설비일상점검일지", "설비일일점검표"}

# 출력(설비 일상 점검표) 시트 열 배치
OUT_HEADER_ROW = 1
OUT_COL = {"no": 1, "공정단계명": 2, "공정특성": 3, "규격": 4}
OUT_LAST_COL = 35  # AI (규격까지 4열 + 1~31일 = 35열)

# 검토표시본에 붙는 주석 마커 — 이 위치 이후 텍스트는 잘라낸다.
ANNOT_MARKERS = ("🔍", "⚠️", "❌", "✅", "[REVIEW", "[MISSING")


def clean_text(v):
    """검토 주석을 잘라내고 내부 공백/개행을 한 칸으로 정리한다."""
    if v is None:
        return ""
    s = str(v).replace("\r", "")
    cut = len(s)
    for marker in ANNOT_MARKERS:
        i = s.find(marker)
        if i != -1 and i < cut:
            cut = i
    return " ".join(s[:cut].split())


def is_blank(v):
    return v is None or str(v).strip() in ("", "-", "–", "—", "－")


def norm(s):
    return re.sub(r"\s+", "", s or "")


def fmt_num(s):
    s = s.strip()
    try:
        f = float(s)
    except (TypeError, ValueError):
        return s
    return str(int(f)) if f.is_integer() else f"{f:g}"


def _is_number(s):
    try:
        float(str(s).strip())
        return True
    except (TypeError, ValueError):
        return False


def format_spec(lower, nominal, upper, unit):
    lo, up, nom = clean_text(lower), clean_text(upper), clean_text(nominal)
    unit_s = clean_text(unit)
    # 단위가 비어있거나 대시일 때만 생략한다. 'Etc' 같은 값은 규격과 함께 그대로 쓴다.
    if unit_s in ("", "-", "–", "—", "－"):
        unit_s = ""

    has_lo, has_up, has_nom = not is_blank(lo), not is_blank(up), not is_blank(nom)

    if has_lo and has_up:
        return f"{fmt_num(lo)} ~ {fmt_num(up)} {unit_s}".strip()
    if has_nom:
        # 숫자형 Nominal이면 단위를 붙이고, 텍스트 규격이면 그대로 둔다
        if unit_s and _is_number(nom):
            return f"{fmt_num(nom)} {unit_s}".strip()
        return nom
    if has_lo:
        return f"{fmt_num(lo)} 이상 {unit_s}".strip()
    if has_up:
        return f"{fmt_num(up)} 이하 {unit_s}".strip()
    return ""


def equipment_sheet_name(raw):
    s = clean_text(raw)
    m = re.search(r"(\d+)\s*호기", s)
    if m:
        return m.group(1)
    s = re.sub(r"^[Mm]", "", s).strip()
    return s or "설비"


def extract_rows(cp_ws):
    """관리기록방법이 설비 일상 점검표인 행을 위에서부터 순서대로 수집한다."""
    rows = []
    current_step = ""
    for r in range(CP_DATA_START_ROW, cp_ws.max_row + 1):
        step_val = cp_ws[f'{CP_COL["공정단계명"]}{r}'].value
        if not is_blank(step_val):
            current_step = clean_text(step_val)

        method = clean_text(cp_ws[f'{CP_COL["관리기록방법"]}{r}'].value)
        if norm(method) not in TARGET_METHODS:
            continue

        char_raw = cp_ws[f'{CP_COL["공정특성"]}{r}'].value
        char = clean_text(char_raw)
        if is_blank(char):
            continue

        spec = format_spec(
            cp_ws[f'{CP_COL["규격하한"]}{r}'].value,
            cp_ws[f'{CP_COL["Norminal"]}{r}'].value,
            cp_ws[f'{CP_COL["규격상한"]}{r}'].value,
            cp_ws[f'{CP_COL["단위"]}{r}'].value,
        )

        needs_review = (
            "[MISSING" in str(char_raw)
            or "[REVIEW" in str(char_raw)
            or spec == ""
        )
        rows.append(
            {
                "공정단계명": current_step,
                "공정특성": char,
                "규격": spec,
                "검토필요": needs_review,
            }
        )
    return rows


def pick_sheet(wb):
    if "1" in wb.sheetnames:
        return wb["1"]
    for sn in wb.sheetnames:
        if norm(wb[sn]["B1"].value) == "공정단계명":
            return wb[sn]
    return wb[wb.sheetnames[-1]]


def style_table(ws, last_row):
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    no_border = Border()
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for r in range(1, last_row + 1):
        for c in range(1, OUT_LAST_COL + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            if r == OUT_HEADER_ROW:
                cell.alignment = center
                cell.font = Font(bold=True)
            else:
                cell.alignment = left if c in (2, 3, 4) else center

    # 마지막 데이터행 아래 / 표 오른쪽(AI 초과)에 템플릿이 남긴 서식 제거
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            if r > last_row or c > OUT_LAST_COL:
                ws.cell(row=r, column=c).border = no_border

    for col, width in {1: 6, 2: 24, 3: 30, 4: 26}.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    for col in range(5, OUT_LAST_COL + 1):
        ws.column_dimensions[get_column_letter(col)].width = 3.6
    for r in range(OUT_HEADER_ROW + 1, last_row + 1):
        ws.row_dimensions[r].height = 30


def write_table(ws, rows):
    # 템플릿에 남아있을 수 있는 기존 본문 제거
    for r in range(OUT_HEADER_ROW + 1, ws.max_row + 1):
        for c in range(1, OUT_LAST_COL + 1):
            ws.cell(row=r, column=c).value = None

    for idx, row in enumerate(rows, start=1):
        r = OUT_HEADER_ROW + idx
        ws.cell(row=r, column=OUT_COL["no"], value=idx)
        ws.cell(row=r, column=OUT_COL["공정단계명"], value=row["공정단계명"])
        ws.cell(row=r, column=OUT_COL["공정특성"], value=row["공정특성"])
        ws.cell(row=r, column=OUT_COL["규격"], value=row["규격"])

    style_table(ws, OUT_HEADER_ROW + len(rows))


def generate(cp_path, template_path, output_path):
    cp_wb = load_cp_workbook(cp_path)
    cp_ws = cp_wb.active

    sheet_name = equipment_sheet_name(cp_ws[CP_HEADER["대상설비"]].value)
    rows = extract_rows(cp_ws)

    wb = openpyxl.load_workbook(template_path, data_only=False)
    ws = pick_sheet(wb)
    keep = ws.title
    for sn in list(wb.sheetnames):
        if sn != keep:
            del wb[sn]

    write_table(ws, rows)
    ws.title = sheet_name
    wb.save(output_path)

    return {
        "설비호기": sheet_name,
        "시트명": sheet_name,
        "항목수": len(rows),
        "CP에서_못찾은_항목": [r["공정특성"] for r in rows if r["검토필요"]],
    }


if __name__ == "__main__":
    cp_path, template_path, output_path = sys.argv[1], sys.argv[2], sys.argv[3]
    print(generate(cp_path, template_path, output_path))
