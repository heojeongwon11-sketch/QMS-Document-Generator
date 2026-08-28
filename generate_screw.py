"""
관리계획서(CP) -> 스크류 점검일지(마모도 검사 기준) 생성기

정책 (사용자 확정 2026-08-27):
    - CP 공정특성(L열)이 '스크류 직경'인 항목의 규격(규격하한~규격상한)을 가져온다.
    - screw_template.xlsx 의 '스크류' 블록에서
        · 판정 OK 치수 칸  = CP 규격 그대로 (예: 25.70~25.99)
        · 판정 NG 치수 칸  = OK 하한값 바로 아래 (하한 - 1 LSB) + ' 이하'
          예) OK 15.65~15.95  ->  NG 15.649 이하
    - 그 외 칸(체크링 치수, 외관, 판정, 조치 등)은 템플릿 그대로 둔다.
    - 시트명 = 설비번호. 예: '135호기' -> '135', 'M135' -> '135'.
"""
import re
import sys

import openpyxl
from openpyxl.utils import get_column_letter

from cp_loader import load_cp_workbook

CP_HEADER = {"대상설비": "M3"}
CP_DATA_START_ROW = 11
CP_COL = {"공정특성": "L", "규격하한": "T", "Norminal": "U", "규격상한": "W", "단위": "X"}

SCREW_CHAR_KEY = "스크류직경"  # 공백 제거 후 비교
ANNOT_MARKERS = ("🔍", "⚠️", "❌", "✅", "[REVIEW", "[MISSING")

# 템플릿에서 스크류 블록을 못 찾을 때 쓰는 기본 좌표
DEFAULT_OK_CELL = "I4"
DEFAULT_NG_CELL = "I6"


def clean_text(v):
    """검토표시본 주석(🔍/⚠️/❌ 등)을 잘라내고 공백/개행을 정리한다."""
    if v is None:
        return ""
    s = str(v).replace("\r", "")
    cut = len(s)
    for marker in ANNOT_MARKERS:
        i = s.find(marker)
        if i != -1 and i < cut:
            cut = i
    return " ".join(s[:cut].split())


def is_blank(v):
    return v is None or str(v).strip() in ("", "-", "–", "—", "－")


def norm(v):
    return re.sub(r"\s+", "", "" if v is None else str(v))


def equipment_sheet_name(raw):
    s = clean_text(raw)
    m = re.search(r"(\d+)\s*호기", s)
    if m:
        return m.group(1)
    s = re.sub(r"^[Mm]", "", s).strip()
    return s or "설비"


def find_screw_spec(cp_ws):
    for r in range(CP_DATA_START_ROW, cp_ws.max_row + 1):
        char = norm(clean_text(cp_ws[f'{CP_COL["공정특성"]}{r}'].value))
        if char.startswith(SCREW_CHAR_KEY):
            return {
                "하한": clean_text(cp_ws[f'{CP_COL["규격하한"]}{r}'].value),
                "nominal": clean_text(cp_ws[f'{CP_COL["Norminal"]}{r}'].value),
                "상한": clean_text(cp_ws[f'{CP_COL["규격상한"]}{r}'].value),
                "행": r,
            }
    return None


def format_ok(spec):
    lo, up, nom = spec["하한"], spec["상한"], spec["nominal"]
    if not is_blank(lo) and not is_blank(up):
        return f"{lo}~{up}"
    if not is_blank(nom):
        return nom
    if not is_blank(lo):
        return f"{lo} 이상"
    if not is_blank(up):
        return f"{up} 이하"
    return ""


def format_ng(spec):
    base = spec["하한"]
    if is_blank(base):
        base = spec["nominal"] if not is_blank(spec["nominal"]) else spec["상한"]
    base = str(base).strip()
    try:
        val = float(base)
    except ValueError:
        return ""
    decimals = len(base.split(".")[-1]) + 1 if "." in base else 3
    ng = round(val - 10 ** (-decimals), decimals)
    return f"{ng:.{decimals}f} 이하"


def find_screw_cells(ws):
    """'스크류' 블록의 '치수' 열과 OK/NG 판정 행을 찾아 (OK칸, NG칸) 좌표를 반환한다."""
    screw_col = screw_row = None
    for row in ws.iter_rows():
        for c in row:
            if norm(c.value) == "스크류":
                screw_col, screw_row = c.column, c.row
    if screw_col is None:
        return DEFAULT_OK_CELL, DEFAULT_NG_CELL

    dim_col = ok_row = ng_row = None
    for row in ws.iter_rows(min_row=screw_row):
        for c in row:
            if c.column < screw_col:
                continue
            v = norm(c.value)
            if v == "치수" and dim_col is None:
                dim_col = c.column
            elif v == "OK" and ok_row is None:
                ok_row = c.row
            elif v == "NG" and ng_row is None:
                ng_row = c.row
    if not (dim_col and ok_row and ng_row):
        return DEFAULT_OK_CELL, DEFAULT_NG_CELL
    col = get_column_letter(dim_col)
    return f"{col}{ok_row}", f"{col}{ng_row}"


def generate(cp_path, template_path, output_path):
    cp_wb = load_cp_workbook(cp_path)
    cp_ws = cp_wb.active

    sheet_name = equipment_sheet_name(cp_ws[CP_HEADER["대상설비"]].value)
    spec = find_screw_spec(cp_ws)
    ok_text = format_ok(spec) if spec else ""
    ng_text = format_ng(spec) if spec else ""

    wb = openpyxl.load_workbook(template_path, data_only=False)
    ws = wb[wb.sheetnames[0]]
    ok_cell, ng_cell = find_screw_cells(ws)
    if ok_text:
        ws[ok_cell] = ok_text
    if ng_text:
        ws[ng_cell] = ng_text
    ws.title = sheet_name
    wb.save(output_path)

    return {
        "설비호기": sheet_name,
        "시트명": sheet_name,
        "스크류직경_OK": ok_text,
        "스크류직경_NG": ng_text,
        "경고": [] if spec else ["CP에서 공정특성 '스크류 직경' 항목을 찾지 못했습니다."],
    }


if __name__ == "__main__":
    cp_path, template_path, output_path = sys.argv[1], sys.argv[2], sys.argv[3]
    print(generate(cp_path, template_path, output_path))
